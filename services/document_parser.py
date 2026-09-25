from __future__ import annotations

import re
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from hashlib import sha256
from io import BytesIO
from pathlib import Path

from charset_normalizer import from_bytes
from docx import Document
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx", ".doc"}


class DocumentParsingError(Exception):
    """Raised when an uploaded document cannot be parsed."""


@dataclass(frozen=True)
class ParsedBlock:
    """A section of extracted text with citation metadata."""

    text: str
    page_number: int | None = None
    paragraph_number: int | None = None
    heading: str | None = None

    @property
    def location_label(self) -> str:
        """Return a readable citation location."""

        location_parts: list[str] = []

        if self.page_number is not None:
            location_parts.append(f"Page {self.page_number}")

        if self.paragraph_number is not None:
            location_parts.append(f"Paragraph {self.paragraph_number}")

        if self.heading:
            location_parts.append(self.heading)

        return " · ".join(location_parts) or "Location unavailable"


@dataclass(frozen=True)
class ParsedDocument:
    """A document and its extracted, traceable text blocks."""

    document_id: str
    filename: str
    extension: str
    content_hash: str
    blocks: tuple[ParsedBlock, ...]

    @property
    def full_text(self) -> str:
        """Return all extracted document text."""

        return "\n\n".join(block.text for block in self.blocks)

    @property
    def word_count(self) -> int:
        """Return an approximate word count."""

        return len(self.full_text.split())

    @property
    def page_count(self) -> int | None:
        """Return the highest PDF page number when available."""

        page_numbers = [
            block.page_number
            for block in self.blocks
            if block.page_number is not None
        ]

        return max(page_numbers) if page_numbers else None


def normalise_text(text: str) -> str:
    """Clean extracted text without destroying paragraph boundaries."""

    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)

    return text.strip()


def split_plain_text_into_blocks(text: str) -> tuple[ParsedBlock, ...]:
    """Split plain text into numbered paragraphs."""

    cleaned_text = normalise_text(text)

    if not cleaned_text:
        return ()

    paragraphs = [
        paragraph.strip()
        for paragraph in re.split(r"\n\s*\n", cleaned_text)
        if paragraph.strip()
    ]

    return tuple(
        ParsedBlock(
            text=paragraph,
            paragraph_number=index,
        )
        for index, paragraph in enumerate(paragraphs, start=1)
    )


def decode_text_file(file_bytes: bytes) -> str:
    """Decode a TXT file while detecting its character encoding."""

    detected_text = from_bytes(file_bytes).best()

    if detected_text is None:
        raise DocumentParsingError(
            "The TXT file encoding could not be identified."
        )

    return str(detected_text)


def parse_pdf(file_bytes: bytes) -> tuple[ParsedBlock, ...]:
    """Extract text page by page from a searchable PDF."""

    try:
        reader = PdfReader(BytesIO(file_bytes))
    except Exception as error:
        raise DocumentParsingError(
            "The PDF could not be opened. It may be damaged or encrypted."
        ) from error

    blocks: list[ParsedBlock] = []

    for page_number, page in enumerate(reader.pages, start=1):
        try:
            page_text = normalise_text(page.extract_text() or "")
        except Exception:
            page_text = ""

        if page_text:
            blocks.append(
                ParsedBlock(
                    text=page_text,
                    page_number=page_number,
                )
            )

    if not blocks:
        raise DocumentParsingError(
            "No searchable text was found in this PDF. "
            "It may be a scanned document that requires OCR."
        )

    return tuple(blocks)


def parse_docx(file_bytes: bytes) -> tuple[ParsedBlock, ...]:
    """Extract paragraphs and headings from a DOCX file."""

    try:
        document = Document(BytesIO(file_bytes))
    except Exception as error:
        raise DocumentParsingError(
            "The DOCX file could not be opened."
        ) from error

    blocks: list[ParsedBlock] = []
    current_heading: str | None = None
    paragraph_number = 0

    for paragraph in document.paragraphs:
        text = normalise_text(paragraph.text)

        if not text:
            continue

        style_name = paragraph.style.name if paragraph.style else ""

        if style_name.lower().startswith("heading"):
            current_heading = text
            continue

        paragraph_number += 1

        blocks.append(
            ParsedBlock(
                text=text,
                paragraph_number=paragraph_number,
                heading=current_heading,
            )
        )

    if not blocks:
        raise DocumentParsingError(
            "No readable paragraphs were found in this DOCX file."
        )

    return tuple(blocks)


def parse_legacy_doc(file_bytes: bytes) -> tuple[ParsedBlock, ...]:
    """Extract text from an older DOC file using antiword."""

    antiword_path = shutil.which("antiword")

    if antiword_path is None:
        raise DocumentParsingError(
            "DOC support requires the antiword system tool. "
            "PDF, TXT and DOCX are already supported."
        )

    temporary_path: str | None = None

    try:
        with tempfile.NamedTemporaryFile(
            suffix=".doc",
            delete=False,
        ) as temporary_file:
            temporary_file.write(file_bytes)
            temporary_path = temporary_file.name

        result = subprocess.run(
            [antiword_path, temporary_path],
            capture_output=True,
            check=False,
            timeout=60,
        )

        if result.returncode != 0:
            raise DocumentParsingError(
                "The DOC file could not be converted to readable text."
            )

        text = decode_text_file(result.stdout)
        blocks = split_plain_text_into_blocks(text)

        if not blocks:
            raise DocumentParsingError(
                "No readable text was found in this DOC file."
            )

        return blocks

    except subprocess.TimeoutExpired as error:
        raise DocumentParsingError(
            "The DOC file took too long to process."
        ) from error

    finally:
        if temporary_path:
            Path(temporary_path).unlink(missing_ok=True)


def parse_document(filename: str, file_bytes: bytes) -> ParsedDocument:
    """Parse a supported document and preserve citation metadata."""

    if not filename:
        raise DocumentParsingError("The uploaded file has no filename.")

    if not file_bytes:
        raise DocumentParsingError(f"{filename} is empty.")

    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentParsingError(
            f"{extension or 'This file type'} is not supported."
        )

    content_hash = sha256(file_bytes).hexdigest()
    document_id = content_hash[:16]

    if extension == ".pdf":
        blocks = parse_pdf(file_bytes)
    elif extension == ".txt":
        text = decode_text_file(file_bytes)
        blocks = split_plain_text_into_blocks(text)
    elif extension == ".docx":
        blocks = parse_docx(file_bytes)
    else:
        blocks = parse_legacy_doc(file_bytes)

    if not blocks:
        raise DocumentParsingError(
            f"No readable text was found in {filename}."
        )

    return ParsedDocument(
        document_id=document_id,
        filename=filename,
        extension=extension,
        content_hash=content_hash,
        blocks=blocks,
    )